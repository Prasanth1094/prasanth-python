from flask import Blueprint, jsonify, request, render_template, send_file
import pandas as pd
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
from docx import Document

from src.myProject.users.user_db import (
    get_users,
    add_user,
    get_user_by_id,
    delete_user,
    update_user,
)

user_routes = Blueprint("user_routes", __name__)


@user_routes.route("/user", methods=["GET"])
def template():
    return render_template("user.html")


@user_routes.route("/users", methods=["GET", "POST"])
def get_user():
    if request.method == "GET":
        rows = get_users()
        data = []
        for row in rows:
            result = {"id": row[0], "name": row[1], "email": row[2], "age": row[3]}
            data.append(result)
        return jsonify(data)
    elif request.method == "POST":
        try:
            add_user(request)
            return jsonify({"message": "User saved successfully"}), 200
        except Exception as e:
            print(e, "exception")
            return (
                jsonify(error=str(e)),
                400,
            )  # 400 is the HTTP status code for "Bad Request"
    else:
        return jsonify({"message": "Method Not Allowed"}), 405


@user_routes.route("/users/<int:id>", methods=["GET", "PUT", "DELETE"])
def get_userbyid(id):
    if request.method == "GET":
        row = get_user_by_id(id)
        data = {
            "id": row["id"],
            "name": row["name"],
            "age": row["age"],
            "email": row["email"],
            # add more keys as needed
        }
        return jsonify(data)
    elif request.method == "PUT":
        try:
            update_user(id, request)
            return jsonify({"message": "User updated successfully"}), 200
        except Exception as e:
            print(e, "exception")
            return (
                jsonify(error=str(e)),
                400,
            )  # 400 is the HTTP status code for "Bad Request"
    elif request.method == "DELETE":
        try:
            delete_user(id)
            return jsonify({"message": "User deleted successfully"}), 200
        except Exception as e:
            print(e, "exception")
            return jsonify(error=str(e)), 400
    else:
        return jsonify({"message": "Method Not Allowed"}), 405


@user_routes.route("/users/reports", methods=["GET"])
def get_reports():
    data = get_users()
    column_names = ["Id", "Name", "Email", "Age"]
    df = pd.DataFrame(data, columns=column_names)
    df.index = df.index + 1
    df.insert(0, "Serio.No", df.index)
    report_file_path = "report.csv"
    df.to_csv(report_file_path, index=False)
    return send_file(report_file_path, as_attachment=True)

@user_routes.route("/users/pdf", methods=["GET"])
def get_reports_as_pdf():
    data = get_users()
    column_names = ["Id", "Name", "Email", "Age"]
    df = pd.DataFrame(data, columns=column_names)
    df.index = df.index + 1
    df.insert(0, "Serio.No", df.index)
    
     # Create a plot
    fig, ax =plt.subplots(figsize=(12,4))
    ax.axis('tight')
    ax.axis('off')
    ax.table(cellText=df.values, colLabels=df.columns, cellLoc = 'center', loc='center')

    # Save the plot as a PDF
    pdf = PdfPages("report.pdf")
    pdf.savefig(fig, bbox_inches='tight')
    pdf.close()

    return send_file("report.pdf", as_attachment=True)

@user_routes.route("/users/docx", methods=["GET"])
def get_reports_as_docx():
    data = get_users()
    column_names = ["Id", "Name", "Email", "Age"]
    df = pd.DataFrame(data, columns=column_names)
    df.index = df.index + 1
    df.insert(0, "Serio.No", df.index)
    
     # Create a new Document
    doc = Document()

    # Add a table to the document
    table = doc.add_table(df.shape[0]+1, df.shape[1])

    # Add the headers
    for j in range(df.shape[-1]):
        table.cell(0,j).text = df.columns[j]

    # Add the rest of the data frame
    for i in range(df.shape[0]):
        for j in range(df.shape[-1]):
            table.cell(i+1,j).text = str(df.values[i,j])

    # Save the doc
    doc.save("report.docx")

    return send_file("report.docx", as_attachment=True)
